# WHWD_9_30_Combined.py
#
# PURPOSE
# -------
# Creates a combined participant-level WHWD workbook where each treatment-
# effectiveness tab contains:
#
#   participantidentifier
#   WHWD Category binary columns
#   WHWD SubCategory 1 binary columns
#   WHWD Subcategory 2 binary columns
#   all variables from the 16 cleaned survey tabs
#
# WHWD is rebuilt from August12thWHWD_python_output.xlsx so the subcategories
# are available. The WHWD tabs already present in 9.30.26_WHWD_SurveyWorkbook.xlsx
# are intentionally ignored.
#
# Missing values in the 16 appended surveys remain blank/NA.
# WHWD presence variables are 0/1.
#
# Header colors identify where each block of variables came from.

import os
import re
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = "/Users/jameshunt/Desktop/BIC Data/Treatment Effectivness"

WHWD_MAPPED_FILE = os.path.join(
    BASE_DIR,
    "August12thWHWD_python_output.xlsx"
)

SURVEY_WORKBOOK = os.path.join(
    BASE_DIR,
    "9.30.26_WHWD_SurveyWorkbook.xlsx"
)

OUTPUT_FILE = os.path.join(
    BASE_DIR,
    "9.30.26_WHWD_Combined_Effectiveness.xlsx"
)

CODED_SHEET = "Treatment_Coded_Python"
LONG_SHEET = "Data_Long_Python"

RESPONSE_TABS = {
    "VeryEffective": "VeryEffectiveTreatments",
    "MildlyEffective": "MildlyEffectiveTreatments",
    "Ineffective": "IneffectiveTreatments",
    "Damaging": "DamagingTreatments",
    "WorkedThenDidnt": "Worked then didn't",
    "SideEffects": "Side effects",
}

# These are the ONLY tabs read from the 9/30 survey workbook.
# Existing WHWD tabs in that workbook are ignored.
SURVEY_SHEETS = [
    "Baseline Illness History",
    "Basic Information",
    "Beighton Questionnaire",
    "COMPASS-31",
    "DSQ",
    "FSS",
    "FUNCAP",
    "Family Health History",
    "Family Medical History (FMH)",
    "Followup Illness History",
    "GAD-7",
    "Karnofsky Performance Scale",
    "Medical Conditions",
    "Mood & Behavior (MBQ)",
    "Personal History",
    "SF-36",
]

# Prefixes keep identically named variables from different surveys distinct.
SURVEY_PREFIXES = {
    "Baseline Illness History": "BIH",
    "Basic Information": "BasicInfo",
    "Beighton Questionnaire": "Beighton",
    "COMPASS-31": "COMPASS31",
    "DSQ": "DSQ",
    "FSS": "FSS",
    "FUNCAP": "FUNCAP",
    "Family Health History": "FHH",
    "Family Medical History (FMH)": "FMH",
    "Followup Illness History": "FIH",
    "GAD-7": "GAD7",
    "Karnofsky Performance Scale": "Karnofsky",
    "Medical Conditions": "MedCond",
    "Mood & Behavior (MBQ)": "MBQ",
    "Personal History": "Personal",
    "SF-36": "SF36",
}

# Header colors only. Data cells are not colored.
BLOCK_COLORS = {
    "Participant ID": "1F4E78",
    "WHWD Category": "5B9BD5",
    "WHWD SubCategory 1": "70AD47",
    "WHWD Subcategory 2": "ED7D31",
    "Baseline Illness History": "4472C4",
    "Basic Information": "A5A5A5",
    "Beighton Questionnaire": "FFC000",
    "COMPASS-31": "5B9BD5",
    "DSQ": "70AD47",
    "FSS": "C55A11",
    "FUNCAP": "7030A0",
    "Family Health History": "00B0F0",
    "Family Medical History (FMH)": "BF9000",
    "Followup Illness History": "264478",
    "GAD-7": "548235",
    "Karnofsky Performance Scale": "8064A2",
    "Medical Conditions": "C65911",
    "Mood & Behavior (MBQ)": "2F75B5",
    "Personal History": "7F6000",
    "SF-36": "375623",
}


def clean_id(series):
    return series.astype("string").str.strip()


def safe_label(value):
    """Convert a WHWD hierarchy label to a clean column label."""
    if pd.isna(value):
        return ""
    value = str(value).strip()
    return value

def r_safe_name(value):
    """Create stable snake_case names that are easy to use in R."""
    value = str(value).strip().lower()
    value = value.replace("&", " and ")
    value = re.sub(r"[^a-z0-9]+", "_", value)
    value = re.sub(r"_+", "_", value).strip("_")
    if value and value[0].isdigit():
        value = "x_" + value
    return value or "unnamed"


def build_whwd_binary(mapped, response_type, level_col, prefix):
    """
    One row per participant represented in this WHWD response type.
    A hierarchy value is 1 if the participant has >=1 mapped treatment
    with that value; otherwise 0.
    """
    response_rows = mapped[
        mapped["response_type"].eq(response_type)
    ].copy()

    participants = (
        response_rows["participantidentifier"]
        .dropna()
        .astype(str)
        .str.strip()
    )
    participants = participants[participants.ne("")].drop_duplicates().sort_values()
    base = pd.DataFrame({"participantidentifier": participants})

    valid = response_rows[
        response_rows["Status"].str.casefold().eq("mapped")
    ][["participantidentifier", level_col]].copy()

    valid[level_col] = valid[level_col].map(safe_label)
    valid = valid[valid[level_col].ne("")].drop_duplicates()

    if valid.empty:
        return base, []

    valid["value"] = 1

    wide = (
        valid.pivot_table(
            index="participantidentifier",
            columns=level_col,
            values="value",
            aggfunc="max",
            fill_value=0
        )
        .reset_index()
    )

    labels = sorted(
        [c for c in wide.columns if c != "participantidentifier"],
        key=lambda x: str(x).casefold()
    )

    # Different original labels can collapse to the same R-safe name
    # (for example punctuation/capitalization variants). Combine those
    # duplicate columns using max so the participant-level indicator
    # remains binary and each final column name is unique.
    rename = {label: f"{prefix.lower()}_{r_safe_name(label)}" for label in labels}
    wide = wide.rename(columns=rename)

    value_cols = [c for c in wide.columns if c != "participantidentifier"]
    if len(value_cols) != len(set(value_cols)):
        values = wide[value_cols].copy()
        values = values.T.groupby(level=0, sort=False).max().T
        wide = pd.concat(
            [wide[["participantidentifier"]].reset_index(drop=True),
             values.reset_index(drop=True)],
            axis=1
        )

    result = base.merge(wide, on="participantidentifier", how="left")

    output_cols = [c for c in wide.columns if c != "participantidentifier"]
    for col in output_cols:
        result[col] = pd.to_numeric(
            result[col], errors="coerce"
        ).fillna(0).astype("int8")

    return result, output_cols


def load_and_prepare_survey(sheet_name):
    df = pd.read_excel(SURVEY_WORKBOOK, sheet_name=sheet_name)

    if "participantidentifier" not in df.columns:
        raise ValueError(f"{sheet_name} is missing participantidentifier.")

    df["participantidentifier"] = clean_id(df["participantidentifier"])
    df = df[
        df["participantidentifier"].notna()
        & df["participantidentifier"].ne("")
    ].copy()

    # The source workbook should already be participant-level. Stop rather than
    # silently multiplying rows if a duplicate ID is encountered.
    dupes = df["participantidentifier"].duplicated(keep=False)
    if dupes.any():
        examples = df.loc[dupes, "participantidentifier"].drop_duplicates().head(10).tolist()
        raise ValueError(
            f"{sheet_name} has duplicate participantidentifier values. "
            f"Examples: {examples}"
        )

    prefix = SURVEY_PREFIXES[sheet_name]

    # SF-36 clarification:
    # The survey PDF and observed scoring direction confirm these are two
    # different social-functioning items. Their numeric values and all derived
    # SF-36 scores remain unchanged; only the final analysis labels are made
    # explicit so the variables cannot be confused in R.
    semantic_names = {}
    if sheet_name == "SF-36":
        semantic_names = {
            "sf36_social_act": "social_interference_extent",
            "sf36_social_interfere": "social_interference_frequency",
        }

    rename = {}
    source_metadata = {}
    for col in df.columns:
        if col == "participantidentifier":
            continue
        analysis_name = semantic_names.get(col, col)
        final_col = f"{r_safe_name(prefix)}_{r_safe_name(analysis_name)}"
        rename[col] = final_col
        # Keep the source field in the Data Dictionary for auditability.
        source_metadata[final_col] = col

    df = df.rename(columns=rename)

    if len(df.columns) != len(set(df.columns)):
        duplicates = pd.Index(df.columns)[pd.Index(df.columns).duplicated()].tolist()
        raise ValueError(
            f"{sheet_name} produced duplicate final column names after renaming: "
            f"{duplicates[:10]}"
        )

    return df, list(rename.values()), source_metadata


def write_dataframe(ws, df, block_ranges):
    # Header
    for c_idx, col in enumerate(df.columns, start=1):
        cell = ws.cell(row=1, column=c_idx, value=col)
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Data
    for r_idx, row in enumerate(df.itertuples(index=False, name=None), start=2):
        for c_idx, value in enumerate(row, start=1):
            if pd.isna(value):
                value = None
            ws.cell(row=r_idx, column=c_idx, value=value)

    # Color each source block in the header.
    for block_name, start_col, end_col in block_ranges:
        fill = PatternFill("solid", fgColor=BLOCK_COLORS[block_name])
        for col_idx in range(start_col, end_col + 1):
            ws.cell(1, col_idx).fill = fill

    thin = Side(style="thin", color="D9E1F2")
    for cell in ws[1]:
        cell.border = Border(bottom=thin)

    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 42

    # Keep a very wide analytic workbook usable without making columns enormous.
    ws.column_dimensions["A"].width = 23
    for col_idx in range(2, ws.max_column + 1):
        header = str(ws.cell(1, col_idx).value or "")
        width = min(max(len(header) + 2, 12), 28)
        ws.column_dimensions[get_column_letter(col_idx)].width = width


print("\nLoading WHWD mapped output...")
coded = pd.read_excel(WHWD_MAPPED_FILE, sheet_name=CODED_SHEET)

required = {
    "participantidentifier",
    "response_type",
    "Category",
    "SubCategory 1",
    "Subcategory 2",
    "Status",
}
missing = required - set(coded.columns)
if missing:
    raise ValueError(
        "Treatment_Coded_Python is missing required columns: "
        + ", ".join(sorted(missing))
    )

for col in [
    "participantidentifier",
    "response_type",
    "Category",
    "SubCategory 1",
    "Subcategory 2",
    "Status",
]:
    coded[col] = coded[col].astype("string").fillna("").str.strip()

# QA: these deprecated Category labels should be gone after the updated
# 9/30 mapping workbook is used to regenerate the WHWD mapped output.
deprecated_categories = {"Mestinon", "Hormonal"}
found_deprecated = sorted(
    set(coded.loc[coded["Category"].isin(deprecated_categories), "Category"])
)
if found_deprecated:
    raise ValueError(
        "Updated WHWD mapping was not fully applied. Deprecated Category labels "
        f"are still present: {found_deprecated}"
    )

# Narrative is never used here.
coded = coded[
    ~coded["response_type"].str.casefold().eq("treatments narrative")
].copy()

print("\nLoading the 16 cleaned survey tabs...")
survey_data = {}
survey_columns = {}
survey_metadata = {}

for i, sheet in enumerate(SURVEY_SHEETS, start=1):
    print(f"  [{i:02d}/16] {sheet}")
    (
        survey_data[sheet],
        survey_columns[sheet],
        survey_metadata[sheet],
    ) = load_and_prepare_survey(sheet)

print("\nBuilding combined effectiveness workbook...")

# Build a data dictionary while creating the workbook.
dictionary_rows = []

wb = Workbook()
# Remove the default sheet after the first real sheet is created.
default_ws = wb.active

for tab_num, (tab_name, response_type) in enumerate(RESPONSE_TABS.items(), start=1):
    print(f"\n[{tab_num}/6] Building {tab_name}...")

    cat_df, cat_cols = build_whwd_binary(
        coded, response_type, "Category", "WHWD_CAT"
    )
    sub1_df, sub1_cols = build_whwd_binary(
        coded, response_type, "SubCategory 1", "WHWD_SUB1"
    )
    sub2_df, sub2_cols = build_whwd_binary(
        coded, response_type, "Subcategory 2", "WHWD_SUB2"
    )

    # Record WHWD variables for the data dictionary.
    for col in cat_cols:
        dictionary_rows.append([tab_name, col, "WHWD", "Category", col.replace("whwd_cat_", "", 1)])
    for col in sub1_cols:
        dictionary_rows.append([tab_name, col, "WHWD", "SubCategory 1", col.replace("whwd_sub1_", "", 1)])
    for col in sub2_cols:
        dictionary_rows.append([tab_name, col, "WHWD", "Subcategory 2", col.replace("whwd_sub2_", "", 1)])

    # The response-type participant cohort is established by the WHWD response
    # itself. Subcategory blocks are joined onto that same cohort.
    combined = cat_df.copy()
    combined = combined.merge(sub1_df, on="participantidentifier", how="left")
    combined = combined.merge(sub2_df, on="participantidentifier", how="left")

    # Subcategory absence for a WHWD respondent is a true 0.
    for col in sub1_cols + sub2_cols:
        combined[col] = pd.to_numeric(
            combined[col], errors="coerce"
        ).fillna(0).astype("int8")

    # Track source blocks for header colors.
    block_ranges = [("Participant ID", 1, 1)]
    current_col = 2

    if cat_cols:
        block_ranges.append(
            ("WHWD Category", current_col, current_col + len(cat_cols) - 1)
        )
        current_col += len(cat_cols)

    if sub1_cols:
        block_ranges.append(
            ("WHWD SubCategory 1", current_col, current_col + len(sub1_cols) - 1)
        )
        current_col += len(sub1_cols)

    if sub2_cols:
        block_ranges.append(
            ("WHWD Subcategory 2", current_col, current_col + len(sub2_cols) - 1)
        )
        current_col += len(sub2_cols)

    # Append all 16 surveys. Missing survey participation remains blank.
    for sheet in SURVEY_SHEETS:
        survey_df = survey_data[sheet]
        cols = survey_columns[sheet]

        combined = combined.merge(
            survey_df,
            on="participantidentifier",
            how="left",
            validate="one_to_one"
        )

        for final_col, original_col in survey_metadata[sheet].items():
            variable_level = "Survey variable"
            if sheet == "SF-36" and original_col in {
                "sf36_social_act",
                "sf36_social_interfere",
            }:
                variable_level = "Survey variable (clarified label)"
            dictionary_rows.append([
                tab_name,
                final_col,
                sheet,
                variable_level,
                original_col,
            ])

        if cols:
            block_ranges.append(
                (sheet, current_col, current_col + len(cols) - 1)
            )
            current_col += len(cols)

    ws = wb.create_sheet(title=tab_name)
    write_dataframe(ws, combined, block_ranges)

    print(
        f"    Participants: {len(combined):,} | "
        f"Columns: {len(combined.columns):,}"
    )

# Remove default empty sheet.
wb.remove(default_ws)

# Add an R-friendly data dictionary.
dd = wb.create_sheet("Data_Dictionary", 0)
dd.append(["effectiveness_tab", "column_name", "source", "variable_level", "original_name"])
dd.append(["ALL", "participantidentifier", "WHWD", "Identifier", "participantidentifier"])
seen = set()
unique_rows = []
for row in dictionary_rows:
    key_tuple = tuple(row)
    if key_tuple not in seen:
        seen.add(key_tuple)
        unique_rows.append(row)
for row in unique_rows:
    dd.append(row)
for cell in dd[1]:
    cell.fill = PatternFill("solid", fgColor="1F1F1F")
    cell.font = Font(color="FFFFFF", bold=True)
dd.freeze_panes = "A2"
dd.auto_filter.ref = dd.dimensions
for col, width in {"A":20, "B":45, "C":35, "D":22, "E":45}.items():
    dd.column_dimensions[col].width = width

# Add a simple key so the color coding is self-documenting.
key = wb.create_sheet("Color_Key", 1)
key.append(["Block", "Meaning"])
key_rows = [
    ("Participant ID", "Participant identifier"),
    ("WHWD Category", "WHWD participant-level Category indicators (0/1)"),
    ("WHWD SubCategory 1", "WHWD participant-level SubCategory 1 indicators (0/1)"),
    ("WHWD Subcategory 2", "WHWD participant-level Subcategory 2 indicators (0/1)"),
]
key_rows.extend((sheet, f"Variables appended from {sheet}") for sheet in SURVEY_SHEETS)

for block, meaning in key_rows:
    key.append([block, meaning])

for row in range(2, key.max_row + 1):
    block = key.cell(row, 1).value
    fill = PatternFill("solid", fgColor=BLOCK_COLORS[block])
    key.cell(row, 1).fill = fill
    key.cell(row, 1).font = Font(color="FFFFFF", bold=True)

for cell in key[1]:
    cell.fill = PatternFill("solid", fgColor="1F1F1F")
    cell.font = Font(color="FFFFFF", bold=True)

key.column_dimensions["A"].width = 34
key.column_dimensions["B"].width = 72
key.freeze_panes = "A2"

print("\nSaving workbook...")
wb.save(OUTPUT_FILE)

print("\nDONE")
print(f"Output saved to:\n{OUTPUT_FILE}")
print("\nWHWD tabs from the 9/30 survey workbook were ignored.")
print("Narrative was excluded.")
print("Category, SubCategory 1, and Subcategory 2 were rebuilt from the mapped WHWD output.")
print("All 16 cleaned survey tabs were left-joined by participantidentifier.")
print("Column names were converted to R-friendly snake_case.")
print("SF-36 raw social-item labels were clarified; derived SF-36 scores were not changed.")
print("A Data_Dictionary tab was added with source and original variable names.")
